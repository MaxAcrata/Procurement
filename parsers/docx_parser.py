"""DOCX parser using python-docx."""

from __future__ import annotations

import logging
from pathlib import Path

from models import DocumentParseResult, Product
from parsers.utils import extract_products_from_text, map_columns_by_keywords, parse_quantity

logger = logging.getLogger(__name__)


def parse_docx(file_path: Path) -> DocumentParseResult:
    """Extract product data from a DOCX file."""
    try:
        import docx
    except ImportError:
        return DocumentParseResult(
            file_path=file_path,
            file_type="docx",
            products=[],
            errors=["python-docx not installed"],
        )

    products: list[Product] = []
    raw_text_parts: list[str] = []
    errors: list[str] = []

    try:
        doc = docx.Document(str(file_path))

        # Try tables first (structured data)
        for table in doc.tables:
            if len(table.rows) < 2:
                continue
            header = [cell.text.lower().strip() for cell in table.rows[0].cells]
            col_map = map_columns_by_keywords(header)
            if not col_map:
                continue

            for row in table.rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                product = _row_to_product(cells, col_map)
                if product:
                    products.append(product)

        # Fallback: regex parsing from paragraphs
        if not products:
            full_text = "\n".join(p.text for p in doc.paragraphs)
            raw_text_parts.append(full_text)
            products = extract_products_from_text(full_text)

    except Exception as e:
        logger.exception("DOCX parse error for %s", file_path)
        errors.append(f"DOCX parse error: {e}")

    return DocumentParseResult(
        file_path=file_path,
        file_type="docx",
        products=products,
        raw_text="\n".join(raw_text_parts),
        errors=errors,
    )


def _row_to_product(row: list[str], col_map: dict[str, int]) -> Product | None:
    try:
        name = row[col_map.get("name", 0)]
        article = row[col_map.get("article", 1)]
        quantity_str = row[col_map.get("quantity", 2)]
        supplier = row[col_map.get("supplier", -1)] if "supplier" in col_map else "UNKNOWN"
        unit = row[col_map.get("unit", -1)] if "unit" in col_map else ""

        quantity = parse_quantity(quantity_str) if quantity_str else 0.0

        if name:
            return Product(
                name=name.strip(),
                article=article.strip() if article else "",
                quantity=quantity,
                supplier=supplier.strip() if supplier and supplier != "UNKNOWN" else "UNKNOWN",
                unit=unit.strip() if unit and unit != "UNKNOWN" else "",
            )
    except (IndexError, ValueError) as e:
        logger.warning("Skipping unparseable row %r: %s", row, e)
    return None
