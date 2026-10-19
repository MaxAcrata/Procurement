"""Excel file generation for purchase orders and reports."""

from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path

from models import ComparisonItem, SupplierGroup, SupplierMatchStatus

logger = logging.getLogger(__name__)

# ── Safe filename helper ─────────────────────────────────────────────────────

_SAFE_NAME_RE = re.compile(r"[^\w\s-]", re.UNICODE)


def _safe_filename(name: str) -> str:
    """Convert a supplier name to a filesystem-safe string."""
    cleaned = _SAFE_NAME_RE.sub("", name).strip()
    return re.sub(r"\s+", "_", cleaned) or "unknown"


# ── Purchase order ───────────────────────────────────────────────────────────

def generate_purchase_order(
    supplier: SupplierGroup,
    output_dir: Path,
) -> Path:
    """Generate an Excel purchase order for a single supplier.

    Creates: <supplier_name>_<dd.mm.yyyy>.xlsx
    Columns: Наименование, Артикул, Кол-во, Ед.изм
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заказ"

    today = datetime.now().strftime("%d.%m.%Y")

    headers = ["Наименование", "Артикул", "Кол-во", "Ед.изм"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, size=11, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, product in enumerate(supplier.products, start=2):
        ws.cell(row=i, column=1, value=product.name)
        ws.cell(row=i, column=2, value=product.article)
        ws.cell(row=i, column=3, value=product.quantity)
        ws.cell(row=i, column=4, value=product.unit)
        # Center-align "Кол-во" and "Ед.изм"
        ws.cell(row=i, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=4).alignment = Alignment(horizontal="center")

    # Auto-width columns based on content (capped at 60 chars)
    for col in range(1, 5):
        max_length = max(
            (len(str(ws.cell(row=row, column=col).value or "")) for row in range(1, len(supplier.products) + 2)),
            default=10,
        )
        col_letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

    safe_name = _safe_filename(supplier.name)
    filename = f"{safe_name}_{today}.xlsx"
    filepath = output_dir / filename
    wb.save(filepath)
    logger.info("Generated purchase order: %s (%d items)", filepath.name, len(supplier.products))
    return filepath


# ── Comparison report ────────────────────────────────────────────────────────

_STATUS_COLORS: dict[SupplierMatchStatus, str] = {
    SupplierMatchStatus.OK: "C6EFCE",
    SupplierMatchStatus.MISMATCH_QUANTITY: "FFEB9C",
    SupplierMatchStatus.MISSING: "FFC7CE",
    SupplierMatchStatus.EXTRA: "BDD7EE",
}

_STATUS_LABELS: dict[SupplierMatchStatus, str] = {
    SupplierMatchStatus.OK: "OK",
    SupplierMatchStatus.MISMATCH_QUANTITY: "Расхождение",
    SupplierMatchStatus.MISSING: "Отсутствует в счете",
    SupplierMatchStatus.EXTRA: "Лишний в счете",
}


def generate_comparison_report(
    results: list[ComparisonItem],
    output_path: Path,
) -> Path:
    """Generate an Excel comparison report."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сравнение"

    ws.merge_cells("A1:F1")
    ws["A1"] = "Отчет сравнения заявки и счета"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["Наименование", "Артикул", "Кол-во (заявка)", "Кол-во (счет)", "Статус", "Разница"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for i, item in enumerate(results, start=4):
        ws.cell(row=i, column=1, value=item.product.name)
        ws.cell(row=i, column=2, value=item.product.article)
        ws.cell(row=i, column=3, value=item.product.quantity)
        ws.cell(row=i, column=4, value=item.invoice_quantity)
        ws.cell(row=i, column=5, value=_STATUS_LABELS.get(item.status, item.status.value))
        ws.cell(row=i, column=6, value=item.diff)

        color = _STATUS_COLORS.get(item.status)
        if color:
            ws.cell(row=i, column=5).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    for col in range(1, 7):
        max_length = max(
            (len(str(ws.cell(row=row, column=col).value or "")) for row in range(3, len(results) + 5)),
            default=10,
        )
        col_letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

    wb.save(output_path)
    logger.info("Generated comparison report: %s (%d items)", output_path.name, len(results))
    return output_path
