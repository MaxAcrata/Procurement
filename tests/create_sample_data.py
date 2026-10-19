"""Create sample test data for the compare command."""

import openpyxl
from pathlib import Path


def create_sample_order():
    """Create a sample order Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Заявка"

    headers = ["Наименование", "Артикул", "Количество"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    data = [
        ("Бумага А4", "PAPER-A4-500", 100),
        ("Ручка шариковая синяя", "PEN-001", 50),
        ("Степлер настольный", "STAP-10", 10),
    ]

    for row, (name, article, qty) in enumerate(data, 2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=article)
        ws.cell(row=row, column=3, value=qty)

    path = Path("output/sample_order.xlsx")
    wb.save(path)
    return path


def create_sample_invoice():
    """Create a sample invoice Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Счет"

    headers = ["Наименование", "Артикул", "Количество"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    data = [
        ("Бумага А4", "PAPER-A4-500", 100),  # OK
        ("Ручка шариковая синяя", "PEN-001", 45),  # Mismatch
        ("Маркер черный", "MARK-BLK", 25),  # Extra in invoice
    ]

    for row, (name, article, qty) in enumerate(data, 2):
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=article)
        ws.cell(row=row, column=3, value=qty)

    path = Path("output/sample_invoice.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    create_sample_order()
    create_sample_invoice()
    print("Sample files created in output/")
